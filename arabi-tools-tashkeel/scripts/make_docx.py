"""Write the film's sample Word document (src/data/فضل-العلم.docx): a short lesson sheet in plain,
undiacritized Arabic. Run with `uv run --with python-docx python scripts/make_docx.py`."""
import pathlib
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
TITLE = "من فضل العلم"
PARAGRAPHS = [
    "عن أبي هريرة رضي الله عنه أن رسول الله صلى الله عليه وسلم قال: من سلك طريقا يلتمس فيه علما سهل الله له به طريقا إلى الجنة. رواه مسلم.",
    "وعنه أن رسول الله صلى الله عليه وسلم قال: إذا مات الإنسان انقطع عنه عمله إلا من ثلاثة: إلا من صدقة جارية، أو علم ينتفع به، أو ولد صالح يدعو له. رواه مسلم.",
]

def rtl(paragraph, size, bold=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph._p.get_or_add_pPr().append(paragraph._p.makeelement(qn("w:bidi"), {}))
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        r = run._r.get_or_add_rPr()
        r.append(r.makeelement(qn("w:rtl"), {}))
        fonts = r.find(qn("w:rFonts"))
        if fonts is None:
            fonts = r.makeelement(qn("w:rFonts"), {})
            r.append(fonts)
        fonts.set(qn("w:cs"), "Traditional Arabic")
        r.append(r.makeelement(qn("w:szCs"), {qn("w:val"): str(size * 2)}))
        if bold:
            r.append(r.makeelement(qn("w:bCs"), {}))

doc = Document()
rtl(doc.add_paragraph(TITLE), 20, bold=True)
for text in PARAGRAPHS:
    rtl(doc.add_paragraph(text), 16)
out = ROOT / "src/data/فضل-العلم.docx"
doc.save(out)
print(out)
